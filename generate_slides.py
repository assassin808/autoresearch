#!/usr/bin/env python3
"""Generate presentation slides for Omni-Modal S3 Audio Plateau research."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import json, os

# Colors
BLACK = RGBColor(0x1a, 0x1a, 0x2e)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0x99, 0x99, 0x99)
BLUE = RGBColor(0x26, 0x5c, 0xb0)
RED = RGBColor(0xc0, 0x39, 0x2b)
GREEN = RGBColor(0x27, 0xae, 0x60)
ORANGE = RGBColor(0xe6, 0x7e, 0x22)
WHITE = RGBColor(0xff, 0xff, 0xff)
BG_LIGHT = RGBColor(0xf5, 0xf5, 0xf5)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_slide():
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    # White background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    return slide

def add_text(slide, left, top, width, height, text, font_size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT, font_name='Calibri'):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = align
    return tf

def add_bullet_text(slide, left, top, width, height, items, font_size=16, color=DARK):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = 'Calibri'
        p.space_after = Pt(6)
        p.level = 0
    return tf

def add_table(slide, left, top, width, height, data, col_widths=None):
    rows, cols = len(data), len(data[0])
    table_shape = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height))
    table = table_shape.table

    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)

    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(12)
                p.font.name = 'Calibri'
                p.font.color.rgb = DARK
                if r == 0:
                    p.font.bold = True
                    p.font.color.rgb = WHITE
            # Header row styling
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = BLUE
            elif r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xf0, 0xf4, 0xf8)
    return table

# ============================================================
# Slide 1: Title
# ============================================================
slide = add_slide()
add_text(slide, 1, 1.5, 11.3, 1.5,
         "Investigating the Audio Loss Plateau\nin Omni-Modal Joint Training",
         font_size=36, bold=True, color=BLACK, align=PP_ALIGN.CENTER)
add_text(slide, 1, 3.5, 11.3, 0.8,
         "A Systematic Diagnostic Study of Mini-Omni S3 Training",
         font_size=20, color=GRAY, align=PP_ALIGN.CENTER)
add_text(slide, 1, 5.0, 11.3, 0.5,
         "29 Experiments  |  15 Diagnostic Signals  |  6 Optimization Methods  |  30k+ Training Steps",
         font_size=14, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
add_text(slide, 1, 6.0, 11.3, 0.5,
         "March 2026  |  Narval Cluster (A100 GPUs)",
         font_size=14, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

# ============================================================
# Slide 2: Mini-Omni Overview
# ============================================================
slide = add_slide()
add_text(slide, 0.8, 0.4, 11.7, 0.7, "Mini-Omni: Model Architecture", font_size=28, bold=True, color=BLACK)

add_bullet_text(slide, 0.8, 1.3, 5.5, 3.0, [
    "Qwen2-0.5B backbone (~500M params)",
    "Whisper encoder + adapter for audio input",
    "SNAC 24kHz codec: 7 codebook streams",
    "  - CB0: semantic (H=6.03 nats)",
    "  - CB1-6: acoustic (H=7.1-7.9 nats)",
    "Shared embedding table: 181,120 tokens",
    "  - Text: rows 0-151,999 (pretrained)",
    "  - Audio: rows 152,000+ (random init!)",
], font_size=15)

add_text(slide, 7.0, 1.3, 5.5, 0.5, "Three-Stage Training", font_size=20, bold=True, color=BLUE)
add_bullet_text(slide, 7.0, 1.9, 5.5, 2.5, [
    "S1: Train Whisper adapter only (freeze LLM)",
    "S2: Train LLM backbone (freeze adapter)",
    "     Text-only tasks: T1T2, A1T2",
    "S3: Unfreeze everything, 4 task types:",
    "     T1T2, T1A2, A1T2, A1A2",
], font_size=15)

add_text(slide, 7.0, 4.5, 5.5, 0.5, "Key Design Choice", font_size=20, bold=True, color=RED)
add_bullet_text(slide, 7.0, 5.0, 5.5, 1.5, [
    "tie_word_embeddings=True",
    "  lm_head = wte (shared parameter)",
    "  Audio embedding updates flow through lm_head",
], font_size=15)

# ============================================================
# Slide 3: The Anomaly (updated)
# ============================================================
slide = add_slide()
add_text(slide, 0.8, 0.4, 11.7, 0.7, "The Anomaly: Audio Loss Plateau in S3", font_size=28, bold=True, color=BLACK)

add_text(slide, 0.8, 1.3, 5.5, 0.5, "At 3,000 Steps (Baseline)", font_size=20, bold=True, color=RED)

data = [
    ["Metric", "Text", "Audio", "Gap"],
    ["Train Loss", "1.44", "37.25", "26x"],
    ["Loss Reduction", "-77%", "-14.5%", "5.3x"],
    ["Val A1A2", "1.07", "49.1", "46x"],
    ["Top-1 Accuracy", "58.3%", "4.8%", "12x"],
    ["GSNR", "~1.0", "0.08", "12.5x"],
]
add_table(slide, 0.8, 1.9, 5.5, 2.5, data)

add_text(slide, 7.0, 1.3, 5.5, 0.5, "After 16,500 Steps (Chain R3)", font_size=20, bold=True, color=GREEN)
data2 = [
    ["Metric", "Value", "vs Baseline"],
    ["Train Audio Loss", "23.85", "-36%"],
    ["Val A1A2", "32.3", "-34%"],
    ["CB0 Top-10", "44.4%", "+114%"],
    ["Audio GSNR", "0.083", "still low"],
    ["Linear Probe L12", "6.1%", "+28%"],
]
add_table(slide, 7.0, 1.9, 5.5, 2.5, data2)

add_text(slide, 0.8, 4.8, 11.7, 1.5,
         "Key Update: The plateau is NOT permanent. Audio loss continues to decrease\n"
         "with more training steps, but at ~12x slower rate than text (predicted by GSNR).",
         font_size=16, color=DARK)

# ============================================================
# Slide 4: Experimental Setup
# ============================================================
slide = add_slide()
add_text(slide, 0.8, 0.4, 11.7, 0.7, "Experimental Setup", font_size=28, bold=True, color=BLACK)

data = [
    ["Batch", "Experiments", "Steps", "eff_batch", "Key Variable"],
    ["Batch 3", "exp1-4", "3,000", "32", "Baseline + initial methods"],
    ["Batch 4", "exp5-10", "3,000", "32", "Fixed methods + new strategies"],
    ["Batch 5", "exp11-13", "10,000", "32", "Longer training + S2 effect"],
    ["Batch 6", "exp15-17", "5k-30k", "32/192", "Batch size + step count"],
    ["Chain", "R1-R3 (+R4-6)", "27,500", "192", "Mini-omni scale"],
]
add_table(slide, 0.8, 1.2, 11.7, 2.8, data)

add_text(slide, 0.8, 4.3, 5.5, 0.5, "15 Diagnostic Signals", font_size=18, bold=True, color=BLUE)
add_bullet_text(slide, 0.8, 4.8, 5.5, 2.5, [
    "D1-D2: Gradient norm ratio + cosine similarity",
    "D4: Basin width (random perturbation probing)",
    "D5-D6: Parameter displacement + embedding rank",
    "D8: Top-k accuracy + normalized loss",
    "D9: Linear probe on hidden states",
    "D10-D11: Per-task val loss + per-module grad norms",
    "D12-D13: CKA representation drift + cosine collapse",
    "D14-D15: GSNR + text subspace projection",
], font_size=13)

add_text(slide, 7.0, 4.3, 5.5, 0.5, "Hardware & Data", font_size=18, bold=True, color=BLUE)
add_bullet_text(slide, 7.0, 4.8, 5.5, 2.5, [
    "Cluster: Narval (Compute Canada), A100 80GB",
    "Data: VoiceAssistant-400K",
    "  470K samples x 4 tasks = 1.88M sequences",
    "  65GB train.pt + 1.4GB val.pt",
    "Default: LR 2e-6 to 2e-5 cosine, warmup 1500",
    "Checkpoint: post-S1 (Qwen2 + adapter, random audio emb)",
], font_size=13)

# ============================================================
# Slide 5: Observation 1 — Gradient Orthogonality
# ============================================================
slide = add_slide()
add_text(slide, 0.8, 0.4, 11.7, 0.7, "Observation 1: Text and Audio Gradients Are Orthogonal", font_size=28, bold=True, color=BLACK)

add_text(slide, 0.8, 1.2, 5.5, 0.5, "D2: Cosine Similarity (cos phi)", font_size=20, bold=True, color=BLUE)
add_bullet_text(slide, 0.8, 1.8, 5.5, 2.0, [
    "cos(phi) = 0.004 +/- 0.02 across ALL experiments",
    "Not conflicting (would be < 0)",
    "Not aligned (would be > 0)",
    "Gradients operate in orthogonal subspaces",
], font_size=15)

add_text(slide, 0.8, 3.8, 5.5, 0.5, "D15: Text Subspace Projection", font_size=20, bold=True, color=BLUE)
add_bullet_text(slide, 0.8, 4.4, 5.5, 2.0, [
    "Audio gradient energy in text subspace: 0.1%",
    "Random baseline: 0.00005%",
    "Weak but non-zero overlap (~2000x random)",
    "Text and audio learn in separate parameter spaces",
], font_size=15)

add_text(slide, 7.0, 1.2, 5.5, 0.5, "Implication", font_size=20, bold=True, color=RED)
add_bullet_text(slide, 7.0, 1.8, 5.5, 4.0, [
    "Root cause: text and audio embeddings are",
    "in completely different subspaces",
    "",
    "Text embeddings: pretrained, structured",
    "Audio embeddings: random N(0, 0.02)",
    "",
    "Backbone features align with text directions",
    "Audio tokens can't access these features",
    "",
    "=> Text doesn't guide audio learning",
    "=> Gradient projection (PCGrad) doesn't help",
    "     (verified: exp3 grad_proj was ineffective)",
], font_size=15)

# ============================================================
# Slide 6: Observation 2 — GSNR
# ============================================================
slide = add_slide()
add_text(slide, 0.8, 0.4, 11.7, 0.7, "Observation 2: Audio GSNR Explains the Slow Learning", font_size=28, bold=True, color=BLACK)

add_text(slide, 0.8, 1.2, 5.5, 0.5, "D14: Gradient Signal-to-Noise Ratio", font_size=20, bold=True, color=BLUE)
add_bullet_text(slide, 0.8, 1.8, 5.5, 2.5, [
    "GSNR = ||E[g]||^2 / Var(g)",
    "",
    "Audio GSNR (B=2):  0.08",
    "Text GSNR (B=2):   ~1.0  (estimated)",
    "",
    "Audio signal is 12.5x weaker than noise",
    "Each gradient step is mostly random walk",
    "Signal emerges only after many steps (sqrt(N))",
], font_size=15)

add_text(slide, 7.0, 1.2, 5.5, 0.5, "Critical Batch Size Derivation", font_size=20, bold=True, color=BLUE)
add_bullet_text(slide, 7.0, 1.8, 5.5, 4.0, [
    "B_crit = tr(Sigma) / ||G||^2",
    "",
    "B_crit_text  = 2   (needs tiny batch)",
    "B_crit_audio = 25  (needs larger batch)",
    "",
    "Optimal total: B* = B_crit_T + sqrt(B_crit_T * B_crit_A)",
    "             = 2 + sqrt(50) = 9",
    "",
    "Optimal sampling: p* = B_crit_T / B* = 22% text",
    "",
    "We used: B=32-192, p=50%",
    "=> Text over-sampled, audio under-sampled",
    "=> eff=192 wastes 81% of compute vs optimal",
], font_size=14)

# ============================================================
# Slide 7: Observation 3 — Two CKA Paths
# ============================================================
slide = add_slide()
add_text(slide, 0.8, 0.4, 11.7, 0.7, "Observation 3: Two Paths to Same Destination", font_size=28, bold=True, color=BLACK)

data = [
    ["Config", "CKA@L12", "Val A1A2", "Backbone", "Where Audio Learns"],
    ["Direct S3 (10k)", "0.036", "44.9", "Completely reshaped", "Full network (output-first)"],
    ["S2->S3 (10k)", "0.990", "44.8", "Nearly frozen", "Layer 0-2 only (emergent adapter)"],
    ["Chain R3 (16k)", "1.000", "32.3", "Completely frozen", "Layer 0-2 + lm_head"],
]
add_table(slide, 0.8, 1.3, 11.7, 2.2, data)

add_text(slide, 0.8, 3.8, 5.5, 0.5, "Direct S3: Output-First Learning", font_size=18, bold=True, color=BLUE)
add_bullet_text(slide, 0.8, 4.3, 5.5, 2.5, [
    "step 1k: layer23=23.4 >> layer0=2.4",
    "Gradient flows loss -> lm_head -> layer23 -> ...",
    "All layers update -> CKA drops to 0.03",
    "Backbone completely reshaped for audio",
], font_size=14)

add_text(slide, 7.0, 3.8, 5.5, 0.5, "S2->S3: Emergent Adapter", font_size=18, bold=True, color=BLUE)
add_bullet_text(slide, 7.0, 4.3, 5.5, 2.5, [
    "step 30k: layer0=45.2, layer2=37.7 >> mid~13",
    "S2 creates wide text basin -> backbone resists change",
    "Audio signal (GSNR=0.08) too weak to push backbone",
    "Layer 0-2 become spontaneous audio adapter",
    "Similar to LoRA: learning only at the edges",
], font_size=14)

# ============================================================
# Slide 8: Observation 4 — Steps >> Batch Size
# ============================================================
slide = add_slide()
add_text(slide, 0.8, 0.4, 11.7, 0.7, "Observation 4: More Steps >> Larger Batch >> S2 Pre-training", font_size=28, bold=True, color=BLACK)

data = [
    ["Experiment", "S2?", "eff_batch", "Steps", "Total Samples", "Val A1A2"],
    ["Baseline (exp2)", "No", "32", "3,000", "96k", "49.1"],
    ["S2->S3 (exp12)", "Yes", "32", "10,000", "320k", "44.8"],
    ["No S2 (exp13)", "No", "32", "10,000", "320k", "44.9"],
    ["S2->S3 eff=192 (exp15)", "Yes", "192", "5,000", "960k", "44.6"],
    ["No S2 eff=192 (exp17)", "No", "192", "5,000", "960k", "43.9"],
    ["S2->S3 30k (exp16)", "Yes", "32", "30,000", "960k", "36.8"],
    ["Chain R3", "Yes", "192", "16,500", "3.2M", "32.3"],
]
add_table(slide, 0.8, 1.2, 11.7, 3.5, data)

add_bullet_text(slide, 0.8, 5.2, 11.7, 2.0, [
    "Same 960k samples: exp16 (30k steps, eff=32) = 36.8  vs  exp15 (5k steps, eff=192) = 44.6",
    "S2 doesn't help val: exp12 (44.8) = exp13 (44.9) at matched steps",
    "Small batch 2.8x more sample-efficient (B_crit theory: 3.8x predicted)",
], font_size=15)

# ============================================================
# Slide 9: Observation 5 — Linear Probe Gap
# ============================================================
slide = add_slide()
add_text(slide, 0.8, 0.4, 11.7, 0.7, "Observation 5: Backbone Memorizes but Doesn't Generalize Audio", font_size=28, bold=True, color=BLACK)

add_text(slide, 0.8, 1.2, 5.5, 0.5, "D9: Linear Probe Train vs Val", font_size=20, bold=True, color=BLUE)
data = [
    ["Layer", "Train Acc", "Val Acc", "Gap"],
    ["L6 (shallow)", "18.6%", "6.2%", "3x"],
    ["L12 (middle)", "57.3%", "6.1%", "9x"],
    ["L18 (deep)", "77.1%", "6.2%", "12x"],
]
add_table(slide, 0.8, 1.8, 5.0, 2.0, data)

add_bullet_text(slide, 0.8, 4.0, 5.5, 2.5, [
    "Backbone encodes 77% audio info at L18",
    "But this info is position-specific, not general",
    "Cannot learn rules because audio embeddings",
    "  are still changing during training",
    "Chicken-and-egg: embeddings need backbone,",
    "  backbone needs stable embeddings",
], font_size=14)

add_text(slide, 7.0, 1.2, 5.5, 0.5, "Audio Learning IS Happening (Slowly)", font_size=20, bold=True, color=GREEN)
data2 = [
    ["Metric", "3k steps", "16.5k steps", "Change"],
    ["CB0 Top-1", "4.8%", "15.9%", "+3.3x"],
    ["CB0 Top-10", "20.7%", "44.4%", "+2.1x"],
    ["D13 Cosine", "0.018", "0.124", "Structuring"],
    ["Probe Val L12", "4.7%", "6.1%", "+28%"],
    ["Val A1A2", "49.1", "32.3", "-34%"],
]
add_table(slide, 7.0, 1.8, 5.5, 2.5, data2)

add_bullet_text(slide, 7.0, 4.5, 5.5, 2.0, [
    "Embeddings forming structure (cosine 0->0.124)",
    "Not collapsing (would be >0.5)",
    "Top-k accuracy continuously improving",
    "Rate consistent with GSNR=0.08 prediction",
], font_size=14)

# ============================================================
# Slide 10: Theoretical Framework
# ============================================================
slide = add_slide()
add_text(slide, 0.8, 0.4, 11.7, 0.7, "Theoretical Framework: Joint Critical Batch Size", font_size=28, bold=True, color=BLACK)

add_bullet_text(slide, 0.8, 1.2, 11.7, 5.5, [
    "Optimization target:  min C = N x B  (total samples to reach target loss)",
    "",
    "Per-modality convergence:  N_x(B_x) = S_min_x * (1 + B_crit_x / B_x)",
    "",
    "Joint constraint (shared params):  N = max(N_text, N_audio)",
    "",
    "With cos(phi) = 0 (orthogonal, verified by D2 + D15):",
    "    Text and audio are independent optimization problems sharing parameters",
    "",
    "Optimal allocation (audio-bottlenecked):",
    "    B* = B_crit_T + sqrt(B_crit_T * B_crit_A) = 2 + sqrt(50) = 9",
    "    p* = B_crit_T / B* = 22% text, 78% audio",
    "",
    "    Audio stays BELOW B_crit (GSNR < 1) at optimum!",
    "    Because: more steps with noisy gradients > fewer steps with clean gradients",
    "",
    "Efficiency at our batch sizes:",
    "    B=32:   70% efficient (close to optimal)",
    "    B=192:  19% efficient (5x waste)",
    "    B=9:   100% efficient (mathematical optimum)",
], font_size=15)

# ============================================================
# Slide 11: Key Insights
# ============================================================
slide = add_slide()
add_text(slide, 0.8, 0.4, 11.7, 0.7, "Key Insights", font_size=28, bold=True, color=BLACK)

insights = [
    ("1. Not Gradient Conflict",
     "cos(phi)=0, energy overlap=0.1%. Text and audio learn independently in orthogonal subspaces."),
    ("2. GSNR is the Root Cause",
     "Audio GSNR=0.08 means 12x more samples needed than text. This is a noise problem, not optimization."),
    ("3. Plateau Breaks with Steps",
     "Val A1A2: 49 -> 45 -> 33 over 16.5k steps. Audio IS learning, just 12x slower."),
    ("4. S2 Doesn't Help Capability",
     "Same val at matched steps. S2 smooths landscape (basin 2.4x wider) but doesn't speed audio learning."),
    ("5. Backbone is NOT the Bottleneck",
     "CKA=0.03 vs CKA=0.99 reach identical val. Audio bottleneck is in embedding space, not backbone."),
    ("6. Emergent Adapter Behavior",
     "With S2, audio gradients concentrate in layer 0-2 (spontaneous LoRA). Without S2, full network reshapes."),
]

y = 1.2
for title, desc in insights:
    add_text(slide, 0.8, y, 2.5, 0.5, title, font_size=15, bold=True, color=BLUE)
    add_text(slide, 3.4, y, 9.5, 0.5, desc, font_size=14, color=DARK)
    y += 0.85

# ============================================================
# Slide 12: Ongoing + Next Steps
# ============================================================
slide = add_slide()
add_text(slide, 0.8, 0.4, 11.7, 0.7, "Ongoing & Next Steps", font_size=28, bold=True, color=BLACK)

add_text(slide, 0.8, 1.2, 5.5, 0.5, "Ongoing: Mini-Omni Scale Chain", font_size=20, bold=True, color=GREEN)
add_bullet_text(slide, 0.8, 1.8, 5.5, 2.5, [
    "R1-R3 complete (global step 16,500)",
    "R4-R6 queued (to 27,500 = 2.9 epochs)",
    "eff_batch=192, continuous cosine LR schedule",
    "Waiting for cluster maintenance (April 7)",
    "",
    "Audio loss trajectory: 29.9 -> 25.6 -> 23.9",
    "Val A1A2 trajectory: 49 -> 35 -> 32",
    "Expect further improvement at 27.5k steps",
], font_size=14)

add_text(slide, 7.0, 1.2, 5.5, 0.5, "Proposed Experiments", font_size=20, bold=True, color=ORANGE)
add_bullet_text(slide, 7.0, 1.8, 5.5, 4.0, [
    "1. Optimal sampling ratio (p=22% text, 78% audio)",
    "   Based on B_crit theory, simple DataLoader change",
    "",
    "2. LR scaling with batch size",
    "   eff=192 + LR=7x should match small batch efficiency",
    "",
    "3. Per-modality effective batch measurement",
    "   Measure GSNR_text directly (currently estimated)",
    "",
    "4. Longer training (50k+ steps)",
    "   Audio is still improving, no convergence yet",
], font_size=14)

# ============================================================
# Save
# ============================================================
output_path = "slides_omni_basin_shaping.pptx"
prs.save(output_path)
print(f"Saved {output_path}")
print(f"Slides: {len(prs.slides)}")
